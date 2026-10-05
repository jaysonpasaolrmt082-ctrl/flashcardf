"""Creates the master extraction workbook and empty analysis-ready CSVs.

No numerical outcome data are entered here. Rows become analysable only when a
reviewer fills them from the full text and sets verified = 'yes'.
"""
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

# ---- analysis-ready long-format files (the R pipeline reads these) ----
COMMON_TAIL = ["species_stratum","histology_confirmed","overlap_group","include_primary",
               "rob_overall","data_origin","source_location","extractor_1","extractor_2",
               "verified","notes"]
LONG = {
 "binary_outcomes.csv": ["study_id","outcome","event_SACC","n_SACC","event_NSACC","n_NSACC"] + COMMON_TAIL,
 "continuous_outcomes.csv": ["study_id","outcome","unit","mean_SACC","sd_SACC","n_SACC","mean_NSACC","sd_NSACC","n_NSACC",
                             "original_format","q1_SACC","median_SACC","q3_SACC","q1_NSACC","median_NSACC","q3_NSACC","transformation"] + COMMON_TAIL,
 "survival_outcomes.csv": ["study_id","outcome","hr","lower95","upper95","adjusted","covariates","hr_priority",
                           "reference_group","n_SACC","n_NSACC","median_follow_up_months"] + COMMON_TAIL,
 "within_sacc_prognostic.csv": ["study_id","factor","comparison","outcome","hr","lower95","upper95","adjusted","covariates",
                                "hr_priority","n_SACC_total","n_events"] + COMMON_TAIL,
 "molecular_evidence.csv": ["study_id","biomarker","alteration_definition","method","positive_SACC","total_SACC",
                            "positive_NSACC","total_NSACC","reported_p","comparator_type"] + COMMON_TAIL,
 "evidence_direction.csv": ["study_id","feature","direction","basis","source_location","verified"],
 "rob_nos.csv": ["study_id","S1_representativeness","S2_selection_nonexposed","S3_ascertainment_exposure","S4_outcome_absent_at_start",
                 "C1_comparability","O1_assessment_outcome","O2_follow_up_length","O3_follow_up_adequacy","total_stars","rating",
                 "assessor_1","assessor_2","consensus_notes"],
 "prisma_counts.csv": ["box","n","verified","notes"],
}
for fn, cols in LONG.items():
    with open(f"data/analysis_ready/{fn}", "w", newline="") as f:
        csv.writer(f).writerow(cols)
with open("data/analysis_ready/prisma_counts.csv","a",newline="") as f:
    w = csv.writer(f)
    for b in ["db_pubmed","db_scopus","db_wos","db_embase","db_cnki","db_wanfang","db_sinomed","other_sources",
              "duplicates_removed","screened","excluded_title_abstract","reports_sought","reports_not_retrieved",
              "full_text_assessed","full_text_excluded","included_qualitative","included_quantitative"]:
        w.writerow([b, "", "no", ""])

# ---- controlled vocabularies ----
OUTCOMES_BIN = ["male_sex","age_over_60","rectal_location","sigmoid_location","left_sided","advanced_stage_III_IV","T3_T4",
    "LN_metastasis","distant_metastasis","vascular_invasion","lymphovascular_invasion","perineural_invasion","tumor_budding",
    "poor_differentiation","mucinous","signet_ring","multiple_primary","synchronous","concomitant_polyps","positive_margin",
    "KRAS_mut","NRAS_mut","BRAF_mut","TP53_mut","p53_IHC_positive","MSI_H","dMMR","CEA_elevated","CA19_9_elevated","recurrence"]

# ---- master wide workbook (user-specified variable list) ----
MASTER = {
 "STUDY IDENTIFICATION": ["study_id","first_author","publication_year","title","journal","PMID","DOI","country","province","city","hospital","department"],
 "DESIGN": ["study_design","retrospective_prospective","recruitment_start","recruitment_end","follow_up_duration","single_center_multicenter"],
 "SCHISTOSOMIASIS": ["species","diagnostic_method","egg_histology_yes_no","history_based_yes_no","serology_yes_no","stool_test_yes_no","PCR_yes_no","definition_verbatim"],
 "POPULATION": ["n_total","n_SACC","n_NSACC","mean_age_SACC","sd_age_SACC","mean_age_NSACC","sd_age_NSACC","male_SACC","female_SACC","male_NSACC","female_NSACC","BMI_SACC","BMI_NSACC"],
 "TUMOR": ["colon_SACC","colon_NSACC","rectal_SACC","rectal_NSACC","right_colon_SACC","right_colon_NSACC","left_colon_SACC","left_colon_NSACC",
           "sigmoid_SACC","sigmoid_NSACC","stage_I_II_SACC","stage_I_II_NSACC","stage_III_IV_SACC","stage_III_IV_NSACC",
           "LN_positive_SACC","LN_negative_SACC","LN_positive_NSACC","LN_negative_NSACC","metastasis_SACC","metastasis_NSACC",
           "vascular_invasion_SACC","vascular_invasion_NSACC","lymphovascular_invasion_SACC","lymphovascular_invasion_NSACC",
           "perineural_invasion_SACC","perineural_invasion_NSACC","tumor_budding_SACC","tumor_budding_NSACC",
           "multiple_primary_SACC","multiple_primary_NSACC","polyps_SACC","polyps_NSACC","positive_margin_SACC","positive_margin_NSACC",
           "poor_diff_SACC","poor_diff_NSACC","mucinous_SACC","mucinous_NSACC","signet_SACC","signet_NSACC","tumor_size_mean_SACC","tumor_size_mean_NSACC"],
 "MOLECULAR": [f"{m}_{s}" for m in ["KRAS","NRAS","BRAF","TP53","APC","PIK3CA","MSI","dMMR","MLH1_loss","MSH2_loss","MSH6_loss","PMS2_loss","beta_catenin","p53_IHC","Ki67_high","cMYC_amp"]
               for s in ["mut_SACC","total_SACC","mut_NSACC","total_NSACC"]],
 "LABORATORY": ["CEA_elev_SACC","CEA_total_SACC","CEA_elev_NSACC","CEA_total_NSACC","CA199_elev_SACC","CA199_total_SACC","CA199_elev_NSACC","CA199_total_NSACC",
                "Hb_mean_SACC","Hb_mean_NSACC","WBC_mean_SACC","WBC_mean_NSACC","PLT_mean_SACC","PLT_mean_NSACC"],
 "SURVIVAL": ["OS_HR","OS_lower95","OS_upper95","OS_adjusted_yes_no","OS_covariates","DFS_HR","DFS_lower95","DFS_upper95","DFS_adjusted_yes_no","DFS_covariates",
              "RFS_HR","RFS_lower95","RFS_upper95","reported_HR_or_reconstructed","HR_reference_group"],
 "QA": ["extractor_1","extractor_2","date_extracted","verified","notes"],
}
inv = list(csv.DictReader(open("data/study_inventory.csv")))

wb = Workbook()
hdr_font = Font(bold=True, color="FFFFFF"); dom_fill = PatternFill("solid", fgColor="1F4E79")
ws = wb.active; ws.title = "README"
for r in [
 ["SACC meta-analysis - master extraction workbook"],
 [""],
 ["Rules"],
 ["1. Enter only values read from the full text (or calculated from reported raw counts). Record table/page in source_location."],
 ["2. data_origin: reported | calculated_from_raw | KM_reconstructed | converted_median_IQR."],
 ["3. Leave a cell blank when not reported. Do not type 0 for 'not reported'."],
 ["4. Set verified = yes only after a second reviewer has checked the value against the source."],
 ["5. Overlapping cohorts: see sheet 'overlap'. One study per overlap group per outcome."],
 ["6. Pre-filled identifiers come from bibliographic indexes (not full texts) and must also be checked."],
]: ws.append(r)
ws["A1"].font = Font(bold=True, size=14)

ws = wb.create_sheet("master_wide")
row1, row2 = [], []
for dom, cols in MASTER.items():
    row1 += [dom] + [""]*(len(cols)-1); row2 += cols
ws.append(row1); ws.append(row2)
for c in ws[1]:
    c.font = hdr_font; c.fill = dom_fill
for c in ws[2]:
    c.font = Font(bold=True); c.alignment = Alignment(text_rotation=90)
col = {n:i for i,n in enumerate(row2)}
for s in inv:
    r = [""]*len(row2)
    r[col["study_id"]] = s["study_id"]; r[col["first_author"]] = s["first_author"]; r[col["publication_year"]] = s["year"]
    r[col["title"]] = s["title"]; r[col["journal"]] = s["journal"]
    r[col["PMID"]] = "" if s["pmid"]=="NR" else s["pmid"]; r[col["DOI"]] = "" if s["doi"]=="NR" else s["doi"]
    r[col["verified"]] = "no"
    ws.append(r)
ws.freeze_panes = "B3"

for fn, cols in LONG.items():
    sh = wb.create_sheet(fn.replace(".csv","")[:31]); sh.append(cols)
    for c in sh[1]: c.font = Font(bold=True)
    if "verified" in cols:
        dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True); sh.add_data_validation(dv)
        L = sh.cell(row=1, column=cols.index("verified")+1).column_letter; dv.add(f"{L}2:{L}2000")
    if "outcome" in cols and fn == "binary_outcomes.csv":
        dv = DataValidation(type="list", formula1='"' + ",".join(OUTCOMES_BIN[:20]) + '"', allow_blank=True)
        sh.add_data_validation(dv); L = sh.cell(row=1, column=2).column_letter; dv.add(f"{L}2:{L}2000")

sh = wb.create_sheet("inventory")
sh.append(list(inv[0].keys())); [sh.append(list(s.values())) for s in inv]
sh = wb.create_sheet("overlap")
for r in csv.reader(open("data/overlap_matrix.csv")): sh.append(r)
sh = wb.create_sheet("vocab_outcomes"); sh.append(["binary_outcome_code"]); [sh.append([o]) for o in OUTCOMES_BIN]
wb.save("data/SACC_master_extraction.xlsx")
print("ok", len(row2), "master columns")
